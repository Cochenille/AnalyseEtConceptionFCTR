"""Génère une fiche de lab Word autoportante à partir d'une description YAML.

Usage :
    python outils/fiches/generer_fiche.py outils/fiches/labs/lab06.yaml
    python outils/fiches/generer_fiche.py outils/fiches/modele.yaml --sortie C:/temp/essai.docx

Sans --sortie, la fiche est écrite dans docs/public/fiches/<fichier>.
Le format de la description est expliqué dans outils/fiches/README.md.
"""

import argparse
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RACINE = Path(__file__).resolve().parents[2]
DOSSIER_FICHES = RACINE / "docs" / "public" / "fiches"

COURS = "582-A05-RI — Analyse et conception Web"
LARGEUR_UTILE = 9.5  # pouces : lettre paysage, marges de 0,75 po

BLEU_ENTETE = "E0E7FF"
GRIS_EXEMPLE = "F3F4F6"
GRIS_TEXTE = RGBColor(0x56, 0x61, 0x6C)
CASE = "☐"


# --- Petits outils de mise en forme -------------------------------------------

def ombrer(cellule, couleur):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    cellule._tc.get_or_add_tcPr().append(shd)


def repeter_entete(rangee):
    trPr = rangee._tr.get_or_add_trPr()
    entete = OxmlElement("w:tblHeader")
    entete.set(qn("w:val"), "true")
    trPr.append(entete)


def hauteur_min(rangee, pouces):
    rangee.height = Inches(pouces)
    rangee.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def ecrire(cellule, texte, gras=False, italique=False, gris=False, taille=None):
    """Remplace le contenu de la cellule. Un « \\n » crée un nouveau paragraphe."""
    lignes = str(texte).split("\n") if texte is not None else [""]
    cellule.paragraphs[0].text = ""
    for i, ligne in enumerate(lignes):
        p = cellule.paragraphs[0] if i == 0 else cellule.add_paragraph()
        run = p.add_run(ligne)
        run.bold = gras
        run.italic = italique
        if gris:
            run.font.color.rgb = GRIS_TEXTE
        if taille:
            run.font.size = Pt(taille)
    return cellule


def entete_colonne(cellule, titre, aide=None):
    """En-tête de tableau : une question en gras, une précision en dessous."""
    ecrire(cellule, titre, gras=True)
    if aide:
        run = cellule.add_paragraph().add_run(aide)
        run.font.size = Pt(9)
        run.font.color.rgb = GRIS_TEXTE
    ombrer(cellule, BLEU_ENTETE)


def nouveau_tableau(doc, largeurs):
    """Tableau à la grille Word, largeurs relatives ramenées à la largeur utile."""
    total = sum(largeurs)
    tableau = doc.add_table(rows=0, cols=len(largeurs))
    tableau.style = "Table Grid"
    tableau.autofit = False
    pouces = [LARGEUR_UTILE * l / total for l in largeurs]
    for colonne, largeur in zip(tableau.columns, pouces):
        colonne.width = Inches(largeur)
    return tableau, pouces


def ajouter_rangee(tableau, pouces):
    rangee = tableau.add_row()
    for cellule, largeur in zip(rangee.cells, pouces):
        cellule.width = Inches(largeur)
    return rangee


def phrase(doc, etiquette, texte, gris=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run(f"{etiquette} ").bold = True
    run = p.add_run(texte)
    if gris:
        run.font.color.rgb = GRIS_TEXTE
    return p


# --- Les blocs de la fiche -----------------------------------------------------

def identification(doc, desc):
    doc.add_heading("Identification", level=1)
    champs = ["Membres de l'équipe", "Numéro d'équipe"] + desc.get("identification", [])
    tableau, pouces = nouveau_tableau(doc, [1, 3])
    for champ in champs:
        rangee = ajouter_rangee(tableau, pouces)
        ombrer(ecrire(rangee.cells[0], champ, gras=True), BLEU_ENTETE)
        hauteur_min(rangee, 0.35)
    doc.add_paragraph()

    nom = f"lab{desc['numero']:02d}-equipeN.docx"
    phrase(doc, "Remise :", f"cette fiche remplie, nommée {nom}, déposée dans l'équipe Teams du cours.")
    if desc.get("regle"):
        phrase(doc, "Règle d'or :", desc["regle"])


def tableau_liste(doc, bloc):
    """Une rangée par élément (E1, E2…) : les colonnes sont des questions."""
    colonnes = bloc["colonnes"]
    prefixe = bloc.get("prefixe")
    largeurs = ([0.6] if prefixe else []) + [c.get("largeur", 2) for c in colonnes]
    tableau, pouces = nouveau_tableau(doc, largeurs)

    rangee = ajouter_rangee(tableau, pouces)
    repeter_entete(rangee)
    cellules = list(rangee.cells)
    if prefixe:
        entete_colonne(cellules.pop(0), "#")
    for cellule, colonne in zip(cellules, colonnes):
        entete_colonne(cellule, colonne["titre"], colonne.get("aide"))

    if bloc.get("exemple"):
        rangee = ajouter_rangee(tableau, pouces)
        valeurs = (["Ex."] if prefixe else []) + bloc["exemple"]
        for cellule, valeur in zip(rangee.cells, valeurs):
            ombrer(ecrire(cellule, valeur, italique=True, gris=True), GRIS_EXEMPLE)

    for i in range(1, bloc.get("rangees", 5) + 1):
        rangee = ajouter_rangee(tableau, pouces)
        hauteur_min(rangee, bloc.get("hauteur", 0.45))
        if prefixe:
            ecrire(rangee.cells[0], f"{prefixe}{i}", gras=True)


def tableau_grille(doc, bloc):
    """Une rangée par critère, une colonne par objet comparé (Outil 1, Outil 2…)."""
    colonnes = bloc["colonnes"]
    avec_exemple = any("exemple" in ligne for ligne in bloc["lignes"])
    largeurs = [2.2] + ([2] if avec_exemple else []) + [2] * len(colonnes)
    tableau, pouces = nouveau_tableau(doc, largeurs)

    rangee = ajouter_rangee(tableau, pouces)
    repeter_entete(rangee)
    entetes = [bloc.get("titre_critere", "Critère")] + (["Exemple"] if avec_exemple else []) + colonnes
    for cellule, titre in zip(rangee.cells, entetes):
        entete_colonne(cellule, titre)

    for ligne in bloc["lignes"]:
        rangee = ajouter_rangee(tableau, pouces)
        hauteur_min(rangee, bloc.get("hauteur", 0.45))
        critere = rangee.cells[0]
        entete_colonne(critere, ligne["titre"], ligne.get("aide"))
        if avec_exemple:
            ombrer(ecrire(rangee.cells[1], ligne.get("exemple", ""), italique=True, gris=True), GRIS_EXEMPLE)


def zone_texte(doc, bloc):
    """Une case libre pour un paragraphe de réponse."""
    tableau, pouces = nouveau_tableau(doc, [1])
    if bloc.get("exemple"):
        cellule = ajouter_rangee(tableau, pouces).cells[0]
        ombrer(ecrire(cellule, f"Exemple : {bloc['exemple']}", italique=True, gris=True), GRIS_EXEMPLE)
    hauteur_min(ajouter_rangee(tableau, pouces), 0.3 * bloc.get("lignes", 4))


BLOCS = {"liste": tableau_liste, "grille": tableau_grille, "texte": zone_texte}


def cases_a_cocher(doc, cases):
    tableau, pouces = nouveau_tableau(doc, [1])
    cellule = ajouter_rangee(tableau, pouces).cells[0]
    ecrire(cellule, "Avant de passer à la suite, cochez :", gras=True)
    for texte in cases:
        cellule.add_paragraph(f"{CASE}  {texte}")


def partie(doc, numero, desc):
    titre = f"Partie {numero} — {desc['titre']}"
    if desc.get("duree"):
        titre += f" ({desc['duree']})"
    doc.add_heading(titre, level=1).paragraph_format.keep_with_next = True

    phrase(doc, "À faire :", desc["a_faire"])
    if desc.get("besoin"):
        phrase(doc, "Vous aurez besoin de :", desc["besoin"], gris=True)

    for bloc in desc.get("blocs", []):
        if bloc.get("consigne"):
            p = doc.add_paragraph(bloc["consigne"])
            p.paragraph_format.keep_with_next = True
        BLOCS[bloc["type"]](doc, bloc)
        doc.add_paragraph()

    if desc.get("cochez"):
        cases_a_cocher(doc, desc["cochez"])


def aide(doc, sujets):
    doc.add_heading("Aide", level=1).paragraph_format.page_break_before = True
    for sujet in sujets:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run(sujet["titre"]).bold = True
        for ligne in sujet.get("points", []):
            doc.add_paragraph(ligne, style="List Bullet")
        if sujet.get("texte"):
            doc.add_paragraph(sujet["texte"])


# --- Assemblage ----------------------------------------------------------------

def generer(desc, sortie):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    for marge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, marge, Inches(0.75))
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_paragraph(f"Lab {desc['numero']:02d} — {desc['titre']}", style="Title")
    doc.add_paragraph(f"{COURS} · {desc.get('contexte', 'Travail formatif en équipe')}")
    if desc.get("objectif"):
        phrase(doc, "Ce soir :", desc["objectif"])

    identification(doc, desc)
    for numero, desc_partie in enumerate(desc["parties"], start=1):
        partie(doc, numero, desc_partie)
    if desc.get("aide"):
        aide(doc, desc["aide"])

    sortie.parent.mkdir(parents=True, exist_ok=True)
    doc.save(sortie)
    return sortie


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("description", type=Path, help="fichier YAML qui décrit la fiche")
    parser.add_argument("--sortie", type=Path, help="chemin du .docx produit (par défaut : docs/public/fiches/)")
    args = parser.parse_args()

    desc = yaml.safe_load(args.description.read_text(encoding="utf-8"))
    sortie = args.sortie or DOSSIER_FICHES / desc["fichier"]
    print(f"Fiche écrite : {generer(desc, sortie)}")


if __name__ == "__main__":
    main()
